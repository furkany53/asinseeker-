import openpyxl
import sys

path = r"C:\Users\onury\Downloads\SalesReport_[2018-01-01_2026-09-18].xlsx"
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
print('Sheet isimleri:', wb.sheetnames)
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f'\n=== {sheet_name} === (satir={ws.max_row}, sutun={ws.max_column})')
    rows = ws.iter_rows(min_row=1, max_row=3, values_only=True)
    for i, row in enumerate(rows):
        print(f'  satir {i}:', row)
